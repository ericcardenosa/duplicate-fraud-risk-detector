# -*- coding: utf-8 -*-
"""Source reader for the Duplicate & Fraud Risk Detector.

Reads the four synthetic X_DEMO_*.xlsx exports (purchase orders, invoices,
projects and cost centres) with openpyxl and normalises dates, amounts and
references. The source files are opened read-only and never modified.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PO_RE = re.compile(r"(?<![A-Za-z0-9])(O\d{6,})(?!\d)", re.I)
PROJ_RE = re.compile(r"\b(XPR26-\d{5})\b", re.I)


def safe_str(val: Any) -> str:
    return "" if val is None else str(val).strip()


def amount(val: Any) -> float:
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return round(float(val), 2)
    s = safe_str(val).replace("€", "").replace("\u00a0", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return 0.0


def date_from_excel(val: Any) -> datetime | None:
    if isinstance(val, datetime):
        return val.replace(hour=0, minute=0, second=0, microsecond=0)
    if isinstance(val, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(val))).replace(hour=0, minute=0, second=0, microsecond=0)
        except (OverflowError, ValueError):
            return None
    if isinstance(val, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(val.strip()[:10], fmt)
            except ValueError:
                pass
    return None


def dt_text(v: datetime | None) -> str:
    return v.strftime("%d/%m/%Y") if v else ""


def cell_safe(val: Any) -> Any:
    """Avoid Excel formula injection from free-text source cells."""
    if isinstance(val, str) and val.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + val
    return val


def norm(s: Any) -> str:
    return re.sub(r"\s+", "", safe_str(s)).lower()


def extract_refs(s: Any) -> list[str]:
    return list(dict.fromkeys(x.upper() for x in PO_RE.findall(safe_str(s))))


def project_id(s: Any) -> str:
    m = PROJ_RE.search(safe_str(s))
    return m.group(1).upper() if m else ""


def input_file(directory: Path, prefix: str) -> Path:
    found = sorted(directory.glob(prefix + "*.xlsx"))
    if len(found) != 1:
        raise ValueError(f"Se esperaba un archivo '{prefix}*.xlsx' en {directory}; encontrados: {len(found)}")
    return found[0]


def read_source(path: Path, columns: int, start_row: int, end_col: str, max_rows: int = 12000) -> list[list[Any]]:
    """Read the first worksheet from start_row, keeping the first `columns` cells of non-empty rows."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(min_row=start_row, max_row=max_rows, values_only=True):
            r = list(r[:columns]) + [None] * max(0, columns - len(r))
            if any(x is not None and x != "" for x in r):
                rows.append(r)
        return rows
    finally:
        wb.close()


@dataclass
class Sources:
    orders: list[dict]
    invoices: list[dict]
    project_ids: set[str]
    cecos: dict[str, str]
    files: dict[str, str]


def load_data(directory: Path) -> Sources:
    paths = {k: input_file(directory, prefix) for k, prefix in {
        "po": "X_DEMO_Pedidos_Compra", "invoice": "X_DEMO_Facturas_Compras",
        "project": "X_DEMO_Proyectos", "ceco": "X_DEMO_CECOs",
    }.items()}
    po_rows = read_source(paths["po"], 21, 6, "U")
    invoice_rows = read_source(paths["invoice"], 21, 6, "U")
    proj_rows = read_source(paths["project"], 9, 6, "I")
    ceco_rows = read_source(paths["ceco"], 1, 2, "A", max_rows=1500)
    orders = []
    for ix, r in enumerate(po_rows, 6):
        po = safe_str(r[0]).upper()
        if not PO_RE.fullmatch(po):
            continue  # Includes synthetic export footers.
        orders.append(dict(po=po, row=ix, date=date_from_excel(r[1]), vendor=safe_str(r[4]),
                           description=safe_str(r[5]), tags=safe_str(r[6]), account=safe_str(r[7]),
                           project=safe_str(r[8]), subtotal=amount(r[14]), total=amount(r[15]),
                           erp_invoiced=amount(r[16]), approval=safe_str(r[19]), state=safe_str(r[20])))
    invoices = []
    for ix, r in enumerate(invoice_rows, 6):
        num = safe_str(r[1])
        if not num.startswith("FX26-"):
            continue
        refs = extract_refs(r[6])
        invoices.append(dict(num=num, row=ix, date=date_from_excel(r[0]), acct_date=date_from_excel(r[3]),
                             due=date_from_excel(r[4]), vendor=safe_str(r[5]), description=safe_str(r[6]),
                             tags=safe_str(r[7]), account=safe_str(r[8]), payment_method=safe_str(r[9]),
                             project=safe_str(r[10]), subtotal=amount(r[11]), iva=amount(r[12]),
                             withholding=amount(r[13]), total=amount(r[16]), paid=amount(r[17]),
                             outstanding=amount(r[18]), state=safe_str(r[19]), payment_date=date_from_excel(r[20]),
                             refs=refs))
    project_ids = {project_id(r[0]) for r in proj_rows if r and project_id(r[0])}
    cecos = {norm(r[0]):safe_str(r[0]) for r in ceco_rows if r and norm(r[0]) not in ("cecos", "cecosobligatorios")}
    if len({p["po"] for p in orders}) != len(orders):
        raise ValueError("Se han detectado números PO duplicados; corrige las fuentes antes de emitir el informe.")
    if not orders or not invoices or not project_ids or not cecos:
        raise ValueError("Archivos sin datos suficientes: verifica los cuatro ficheros de entrada.")
    return Sources(orders, invoices, project_ids, cecos, {k:v.name for k,v in paths.items()})
