#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DUPLICATE & FRAUD RISK DETECTOR | SOCIEDAD FICTICIA X, S.L. | v2.1.0

Detecta INDICIOS de duplicidad, no fraude ni pagos duplicados probados.
Trabaja exclusivamente con los cuatro ficheros X_DEMO*.xlsx sintéticos.

Uso:
    python fraud_detector.py --input-dir inputs --output output/DUPLICATE_RISK_2026.xlsx
    python fraud_detector.py --days 7 --min-similarity 0.86 --threshold 1.00
    python -m unittest discover -s tests -v

Dependencia: openpyxl (pip install -r requirements.txt).
"""
from __future__ import annotations

import argparse
import collections
import csv
import difflib
import hashlib
import json
import logging
import re
from datetime import datetime
from itertools import combinations
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.source_reader import load_data, amount, norm, safe_str, dt_text, cell_safe

VERSION = "2.1.0"
ENTITY = "SOCIEDAD FICTICIA X, S.L."
NAVY = "#112239"
INK = "#203753"
BLUE = "#275F94"
TEAL = "#087F84"
WHITE = "#FFFFFF"
PAPER = "#F1F5FA"
SUBTLE = "#E4EDF5"
GREEN = "#15745E"
AMBER = "#C18418"
RED = "#B84449"
PURPLE = "#6F52A1"
MUTED = "#62748A"
EUR = '#,##0.00 "€";[Red](#,##0.00 "€");–'
INTEGER = '#,##0'
MONTHS = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
SEVERITY = {"Crítica":0,"Alta":1,"Media":2,"Baja":3}
RULES = {
    "R01":("Nº factura repetido",95,"Crítica", "Mismo proveedor y número de factura en dos filas."),
    "R02":("Huella documental idéntica",92,"Crítica","Mismo proveedor, fecha, importe y descripción normalizada."),
    "R03":("Importe y fecha idénticos",74,"Alta","Mismo proveedor, fecha e importe, con distinta referencia documental."),
    "R04":("Importe y texto similares",64,"Media","Mismo proveedor e importe; fechas próximas y descripción muy similar."),
    "R05":("Misma PO e importe",69,"Media","Dos facturas de igual importe mencionan la misma PO; puede ser legítimo."),
    "R06":("Facturación supera PO",78,"Alta","Suma de facturas explícitamente vinculadas supera importe total del pedido."),
    "R07":("Proveedor distinto en PO",62,"Media","Proveedor de factura y pedido difieren; verificar datos y autorización."),
}


def clean_desc(value:Any)->str:
    text=safe_str(value).lower()
    # Preservar información documental relevante; eliminar espacios/acentos no intencionadamente.
    return re.sub(r"\s+", " ", re.sub(r"[^\wáéíóúüñ]+", " ", text)).strip()


def sha_id(kind:str, reference:str)->str:
    key=(kind+"|"+reference).encode("utf-8")
    return "CAS-"+hashlib.sha256(key).hexdigest()[:11].upper()


def same_sign(a:float,b:float)->bool:
    return (a>=0 and b>=0) or (a<0 and b<0)


def pair_key(a:dict,b:dict)->tuple[int,int]:
    return tuple(sorted((a["row"], b["row"])))


def risk_band(score:int)->str:
    if score >= 90: return "Crítica"
    if score >= 72: return "Alta"
    if score >= 50: return "Media"
    return "Baja"


def detect(src, days=7, similarity=0.86, threshold=1.0, review:dict|None=None)->dict:
    """Crear casos por PAR de documentos y casos por PO; no sumar exposiciones como fraude.

    Los datos permanecen en memoria, sin alterar ninguna de las cuatro fuentes.
    """
    if days<0 or not 0<=similarity<=1 or threshold<0:
        raise ValueError("Parámetros fuera de rango")
    review=review or {}
    invoices=src.invoices
    po_by_id={o["po"]:o for o in src.orders}
    buckets=collections.defaultdict(list)
    by_id=collections.defaultdict(list)
    by_fingerprint=collections.defaultdict(list)
    by_day_amount=collections.defaultdict(list)
    by_amount=collections.defaultdict(list)
    by_po=collections.defaultdict(list)
    doc_po=collections.defaultdict(list)
    for f in invoices:
        vendor=norm(f["vendor"])
        amt=round(f["total"],2)
        if not vendor:continue
        by_id[(vendor,norm(f["num"]))].append(f)
        if f["date"]:
            by_fingerprint[(vendor,amt,f["date"].date(),clean_desc(f["description"]))].append(f)
            by_day_amount[(vendor,amt,f["date"].date())].append(f)
        by_amount[(vendor,amt)].append(f)
        for ref in f["refs"]:
            if ref in po_by_id:
                by_po[ref].append(f)
                doc_po[f["row"]].append(ref)

    pair_rules=collections.defaultdict(set)
    pair_docs={}
    def add_pairs(group, code, condition=None):
        if len(group)<2:return
        for a,b in combinations(sorted(group,key=lambda x:x["row"]),2):
            if not same_sign(a["total"],b["total"]):continue
            if condition and not condition(a,b):continue
            key=pair_key(a,b)
            pair_rules[key].add(code)
            pair_docs[key]=(a,b)

    for group in by_id.values(): add_pairs(group,"R01")
    for group in by_fingerprint.values():add_pairs(group,"R02")
    for group in by_day_amount.values():add_pairs(group,"R03")
    # Bloqueo proveedor+importe: evita comparación N² entre todo el universo.
    for group in by_amount.values():
        if len(group)<2:continue
        group=sorted([f for f in group if f["date"]],key=lambda f:(f["date"],f["row"]))
        for i,a in enumerate(group):
            for b in group[i+1:]:
                gap=(b["date"]-a["date"]).days
                if gap>days:break
                if gap==0:continue # R03 ya cubre el mismo día
                sim=difflib.SequenceMatcher(None,clean_desc(a["description"]),clean_desc(b["description"])).ratio()
                if sim>=similarity: add_pairs((a,b),"R04")
    for group in by_po.values():
        add_pairs(group,"R05",lambda a,b:abs(a["total"]-b["total"])<=threshold and a["row"]!=b["row"])

    cases=[]
    pair_details=[]
    def date_for_case(a,b):return max((x for x in (a.get("date"),b.get("date")) if x),default=None)
    for key,rule_set in sorted(pair_rules.items()):
        a,b=pair_docs[key]
        primary=sorted(rule_set,key=lambda code:-RULES[code][1])[0]
        points=RULES[primary][1]
        shared=set(doc_po[a["row"]]) & set(doc_po[b["row"]])
        if shared and primary not in ("R01","R02"):points=min(99,points+4)
        # Importe pagado informado ≠ evidencia bancaria. Se utiliza SOLO como contexto.
        if a["paid"]!=0 and b["paid"]!=0:points=min(99,points+2)
        confidence=risk_band(points)
        gap=abs((a["date"]-b["date"]).days) if a["date"] and b["date"] else None
        sim=round(difflib.SequenceMatcher(None,clean_desc(a["description"]),clean_desc(b["description"])).ratio(),3)
        cid=sha_id("PAIR",f"{key[0]}:{key[1]}")
        info=review.get(cid,{})
        cases.append(dict(id=cid,date=date_for_case(a,b),type="PAR",rule=primary,rules=sorted(rule_set),
                          severity=confidence,score=points,vendor=a["vendor"],doc_a=a["num"],doc_b=b["num"],
                          po=", ".join(sorted(shared)),amount=round(min(abs(a["total"]),abs(b["total"])),2),
                          reason=RULES[primary][3],status=info.get("estado","Pendiente"),
                          owner=info.get("responsable",""),notes=info.get("notas",""),evidence=info.get("evidencia",""),
                          source_a=a["row"],source_b=b["row"],gap=gap,similarity=sim,
                          paid_a=a["paid"],paid_b=b["paid"]))
        pair_details.append([cid, a["num"],b["num"],a["vendor"],dt_text(a["date"]),dt_text(b["date"]),
                             a["total"],b["total"], gap if gap is not None else "",sim,", ".join(sorted(rule_set)),
                             a["description"],b["description"],", ".join(sorted(shared)),
                             a["paid"],b["paid"],a["row"],b["row"]])

    overbill=[]
    po_detail=[]
    for p in src.orders:
        fs=by_po.get(p["po"],[])
        # Si una factura cita varias POs, no atribuimos su importe íntegro a cada pedido.
        # Se prorratea para el control con una nota explícita de limitación.
        allocated=round(sum(f["total"] / max(1,len(doc_po[f["row"]])) for f in fs),2)
        diff=round(allocated-p["total"],2)
        over= bool(fs and diff>threshold and p["total"]>0)
        count_multi=sum(len(doc_po[f["row"]])>1 for f in fs)
        flag=("Revisar exceso" if over else ("Sin vínculo" if not fs else "Sin exceso positivo"))
        po_detail.append([p["po"],dt_text(p["date"]),p["vendor"],p["total"],len(fs),allocated,diff,
                          p["erp_invoiced"],count_multi,flag,p["approval"],p["state"],p["row"]])
        if not over:continue
        score= min(94,78+(10 if diff>0.25*p["total"] else 0))
        cid=sha_id("PO",p["po"])
        info=review.get(cid,{})
        cases.append(dict(id=cid,date=p["date"],type="PO",rule="R06",rules=["R06"],
                          severity=risk_band(score),score=score,vendor=p["vendor"],doc_a=p["po"],
                          doc_b=", ".join(f["num"] for f in fs[:5]),po=p["po"],amount=diff,
                          reason="Facturación vinculada mayor que PO. Revisar suplementos, abonos y autorizaciones; importe orientativo.",
                          status=info.get("estado","Pendiente"),owner=info.get("responsable",""),
                          notes=info.get("notas",""),evidence=info.get("evidencia",""),
                          source_a=p["row"],source_b=0,gap=None,similarity=None,paid_a=0,paid_b=0))
        overbill.append([cid,p["po"],p["vendor"],dt_text(p["date"]),p["total"],len(fs),allocated,diff,
                         p["erp_invoiced"],", ".join(f["num"] for f in fs),count_multi,p["approval"],p["row"]])

    vendor_mismatch=[]
    for f in invoices:
        for po in doc_po[f["row"]]:
            p=po_by_id[po]
            if norm(p["vendor"])==norm(f["vendor"]):continue
            cid=sha_id("VENDOR",f'{f["row"]}:{po}')
            info=review.get(cid,{})
            cases.append(dict(id=cid,date=f["date"],type="VINCULO",rule="R07",rules=["R07"],
                              severity="Media",score=62,vendor=f["vendor"],doc_a=f["num"],doc_b=po,po=po,
                              amount=abs(f["total"]),reason="Proveedor de factura distinto del proveedor de PO: puede ser tercero facturador.",
                              status=info.get("estado","Pendiente"),owner=info.get("responsable",""),
                              notes=info.get("notas",""),evidence=info.get("evidencia",""),
                              source_a=f["row"],source_b=p["row"],gap=None,similarity=None,paid_a=f["paid"],paid_b=0))
            vendor_mismatch.append([cid,f["num"],po,f["vendor"],p["vendor"],f["total"],f["row"],p["row"]])

    cases.sort(key=lambda x:(-x["score"],x["rule"],x["id"]))
    if len({x["id"] for x in cases})!=len(cases):raise AssertionError("IDs de caso duplicados")
    invoice_cases=collections.defaultdict(list)
    for c in cases:
        if c["type"]=="PAR":
            invoice_cases[c["source_a"]].append(c)
            invoice_cases[c["source_b"]].append(c)
        elif c["type"]=="VINCULO":invoice_cases[c["source_a"]].append(c)
    invoice_detail=[]
    for f in invoices:
        linked=invoice_cases[f["row"]]
        max_risk=max((c["score"] for c in linked),default=0)
        invoice_detail.append([f["num"],dt_text(f["date"]),dt_text(f["acct_date"]),f["vendor"],
                               f["total"],f["paid"],f["outstanding"],f["state"],", ".join(doc_po[f["row"]]),
                               f["project"],f["account"],f["description"],len(linked),max_risk,
                               ", ".join(c["id"] for c in linked[:5]),f["row"]])
    vendor_stats=collections.defaultdict(lambda:dict(invoice_count=0, volume=0., cases=0, high=0, exposure=0.))
    for f in invoices:
        s=vendor_stats[f["vendor"]];s["invoice_count"]+=1;s["volume"]+=abs(f["total"])
    for c in cases:
        s=vendor_stats[c["vendor"]];s["cases"]+=1;s["high"]+=c["score"]>=72;s["exposure"]+=c["amount"]
    suppliers=sorted([[k,v["invoice_count"],round(v["volume"],2),v["cases"],v["high"],round(v["exposure"],2)]
                      for k,v in vendor_stats.items()],key=lambda x:(-x[3],-x[4],-x[2]))
    monthly=[]
    for month in range(1,13):
        inv=[f for f in invoices if f["date"] and f["date"].year==2026 and f["date"].month==month]
        monthcases=[c for c in cases if c["date"] and c["date"].year==2026 and c["date"].month==month]
        monthly.append([MONTHS[month-1],len(inv),len(monthcases),sum(c["score"]>=72 for c in monthcases),
                        sum(c["rule"]=="R06" for c in monthcases),round(sum(abs(f["total"]) for f in inv),2)])
    controls=[]
    for code,(title,score,severity,description) in RULES.items():
        subset=[c for c in cases if c["rule"]==code]
        controls.append([code,title,score,severity,description,len(subset),
                         round(sum(c["amount"] for c in subset),2),
                         "Duplicidad potencial" if code in ("R01","R02","R03","R04","R05") else "Conciliación / control de compras"])
    return dict(src=src,cases=cases,pairs=pair_details,overbill=overbill,vendor_mismatch=vendor_mismatch,
                invoice_detail=invoice_detail,po_detail=po_detail,suppliers=suppliers,monthly=monthly,
                controls=controls,params=dict(days=days,similarity=similarity,threshold=threshold),
                invoice_case_count=len(set(c["source_a"] for c in cases if c["type"]=="PAR") |
                                       set(c["source_b"] for c in cases if c["type"]=="PAR")))


def read_review_csv(path:Path)->dict:
    if not path.exists():return {}
    out={}
    with path.open(newline="",encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            cid=safe_str(row.get("case_id"))
            if not cid:continue
            state=safe_str(row.get("estado")) or "Pendiente"
            allowed={"Pendiente","En investigación","Justificada","Confirmada","Descartada","Corregida"}
            out[cid]={"estado":state if state in allowed else "Pendiente",
                      "responsable":safe_str(row.get("responsable")),"notas":safe_str(row.get("notas")),
                      "evidencia":safe_str(row.get("evidencia"))}
    return out


def xlcol(n: int) -> str:
    return get_column_letter(n)


def _fill(c: str) -> PatternFill:
    return PatternFill("solid", start_color=c.lstrip("#"), end_color=c.lstrip("#"))


def _font(size=10, bold=False, color=INK, name="Aptos", italic=False) -> Font:
    return Font(name=name, size=size, bold=bold, italic=italic, color=color.lstrip("#"))


def _style(ws, ref, fill=None, font=None, align=None):
    for row in ws[ref]:
        for c in row:
            if fill:
                c.fill = _fill(fill)
            if font:
                c.font = font
            if align:
                c.alignment = align


def sheet_header(ws, name, sub, last_col):
    end = xlcol(last_col)
    ws.merge_cells(f"A1:{end}2")
    _style(ws, f"A1:{end}2", NAVY, _font(18, True, WHITE, "Aptos Display"), Alignment(vertical="center", indent=1))
    ws["A1"] = name
    ws.merge_cells(f"A3:{end}3")
    _style(ws, f"A3:{end}3", INK, _font(10, color="#CAE0F1"), Alignment(vertical="center", indent=1))
    ws["A3"] = sub
    for r, h in ((1, 24), (2, 24), (3, 25)):
        ws.row_dimensions[r].height = h


_TABLES: set[str] = set()
STATUS_LIST = '"Pendiente,En investigación,Justificada,Confirmada,Descartada,Corregida"'


def table_sheet(wb, name, title, sub, headers, rows, widths=None, money=(), score_col=None, status_col=None):
    s = wb.create_sheet(name)
    n = len(headers)
    end = xlcol(n)
    sheet_header(s, title, sub, n)
    for j, h in enumerate(headers, 1):
        c = s.cell(row=5, column=j, value=h)
        c.fill = _fill(BLUE)
        c.font = _font(10, True, WHITE)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    s.row_dimensions[5].height = 39
    last = 5
    if rows:
        assert all(len(row) == n for row in rows), (name, n, len(rows[0]))
        last = 5 + len(rows)
        for i, r in enumerate(rows, 6):
            for j, v in enumerate(r, 1):
                s.cell(row=i, column=j, value=cell_safe(v)).font = _font(10)
    for i in range(1, n + 1):
        s.column_dimensions[xlcol(i)].width = widths[i - 1] if widths else 20
    for i in money:
        for r in range(6, last + 1):
            s.cell(row=r, column=i).number_format = EUR
    if score_col and rows:
        c = xlcol(score_col)
        s.conditional_formatting.add(f"{c}6:{c}{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=TEAL.lstrip("#")))
    if status_col and rows:
        c = xlcol(status_col)
        dv = DataValidation(type="list", formula1=STATUS_LIST, allow_blank=True)
        s.add_data_validation(dv)
        dv.add(f"{c}6:{c}{last}")
    s.freeze_panes = "A6"
    if rows:
        tname = "Tbl" + re.sub(r"[^A-Za-z0-9]", "", name)[:22]
        while tname in _TABLES:
            tname += "x"
        _TABLES.add(tname)
        t = Table(displayName=tname, ref=f"A5:{end}{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        s.add_table(t)
    return s


def card(ws, label, value, left, right):
    ws.merge_cells(f"{left}5:{right}5")
    ws.merge_cells(f"{left}6:{right}7")
    _style(ws, f"{left}5:{right}5", PAPER, _font(10, True, MUTED), Alignment(vertical="center", indent=1))
    _style(ws, f"{left}6:{right}7", WHITE, _font(22, True, NAVY, "Aptos Display"), Alignment(horizontal="right", vertical="center", indent=1))
    ws[f"{left}5"] = label
    ws[f"{left}6"] = value


def build_excel(result: dict, output: Path, source_dir: Path | None = None) -> dict:
    wb = Workbook()
    wb.remove(wb.active)
    _TABLES.clear()
    cases = result["cases"]
    crit = sum(c["severity"] == "Crítica" for c in cases)
    high = sum(c["score"] >= 72 for c in cases)
    pair = sum(c["type"] == "PAR" for c in cases)
    po_cases = sum(c["type"] == "PO" for c in cases)
    total_invoice = round(sum(f["total"] for f in result["src"].invoices), 2)

    case_rows = []
    for c in cases:
        case_rows.append([c['id'], dt_text(c['date']), c['severity'], c['score'], c['rule'], RULES[c['rule']][0],
                          c['type'], c['vendor'], c['doc_a'], c['doc_b'], c['po'], c['amount'], c['reason'],
                          c['status'], c['owner'], c['evidence'], c['notes'], c['source_a'], c['source_b']])
    cs = table_sheet(wb, "02_Casos", "REGISTRO CENTRAL DE CASOS / WORKPAPERS",
                     "Una fila por caso, identificador estable, prioridad y revisión manual. Exposición NO equivale a pérdida.",
                     ["ID caso", "Fecha", "Prioridad", "Score / 100", "Control", "Descripción del control", "Tipo",
                      "Proveedor", "Documento A", "Documento B", "PO", "Exposición orientativa €", "Motivo / evidencia a revisar",
                      "Estado revisión", "Responsable", "Evidencia / enlace interno", "Conclusión / notas", "Fila origen A", "Fila origen B"],
                     case_rows, [22, 15, 14, 17, 13, 32, 14, 36, 21, 35, 18, 25, 83, 22, 25, 45, 56, 16, 16], (12,), 4, 14)
    if case_rows:
        rng = f"C6:C{5 + len(cases)}"
        cs.conditional_formatting.add(rng, FormulaRule(formula=['C6="Crítica"'], fill=_fill("#FBE3E6"), font=Font(color=RED.lstrip("#"), bold=True)))
        cs.conditional_formatting.add(rng, FormulaRule(formula=['C6="Alta"'], fill=_fill("#FFF2D6"), font=Font(color=AMBER.lstrip("#"), bold=True)))
    table_sheet(wb, "03_Pares", "FACTURAS EMPAREJADAS / DETALLE DE COINCIDENCIAS",
                "Compara ambos documentos y sus importes. Una coincidencia NO constituye prueba de duplicidad.",
                ["ID caso", "Factura A", "Factura B", "Proveedor", "Fecha A", "Fecha B", "Total A €", "Total B €",
                 "Días diferencia", "Similitud 0-1", "Reglas activadas", "Descripción A", "Descripción B", "PO compartida",
                 "Pagado informado A €", "Pagado informado B €", "Fila origen A", "Fila origen B"],
                result['pairs'], [21, 22, 22, 34, 16, 16, 21, 21, 18, 17, 22, 60, 60, 24, 24, 24, 17, 17], (7, 8, 15, 16))
    table_sheet(wb, "04_Sobrefacturacion", "FACTURACIÓN VINCULADA QUE SUPERA EL PEDIDO",
                "Comparativa bruto con IVA. Facturas multirreferencia prorrateadas en partes iguales. Verificar cambios de PO.",
                ["ID caso", "PO", "Proveedor", "Fecha PO", "Total PO €", "Facturas", "Facturado atribuido €",
                 "Exceso €", "Facturado ERP €", "Facturas relacionadas", "Docs multirreferencia", "Aprobación", "Fila PO"],
                result['overbill'], [21, 21, 37, 16, 23, 17, 24, 21, 23, 62, 22, 23, 15], (5, 7, 8, 9))
    table_sheet(wb, "05_Facturas", "UNIVERSO COMPLETO DE FACTURAS",
                "Fuente sintética, con referencias y filas de origen; los registros sin alerta permanecen visibles.",
                ["Nº factura", "Fecha emisión", "Fecha contable", "Proveedor", "Total €", "Pagado informado €",
                 "Pendiente informado €", "Estado ERP", "PO(s)", "Proyecto", "Cuenta", "Descripción",
                 "Nº casos par/enlace", "Score máximo", "IDs de caso", "Fila origen"],
                result['invoice_detail'], [24, 19, 19, 38, 20, 22, 24, 19, 25, 38, 22, 67, 20, 19, 55, 15], (5, 6, 7), 14)
    table_sheet(wb, "06_Pedidos", "UNIVERSO COMPLETO DE PEDIDOS",
                "Cruce explícito con facturas por referencia O######; pedidos sin vínculo no se consideran fraudes.",
                ["PO", "Fecha", "Proveedor", "Total PO €", "Nº facturas", "Facturado atribuible €",
                 "Exceso / diferencia €", "Facturado ERP €", "Multirreferencias", "Revisión", "Aprobación", "Estado", "Fila origen"],
                result['po_detail'], [21, 17, 36, 20, 16, 23, 24, 20, 20, 26, 21, 20, 15], (4, 6, 7, 8))
    table_sheet(wb, "07_Proveedores", "ANÁLISIS DE PROVEEDORES / CONCENTRACIÓN",
                "Volumen = importe absoluto para evitar compensación de abonos. Exposición por casos NO aditiva.",
                ["Proveedor", "Nº facturas", "Volumen bruto absoluto €", "Casos totales", "Alta + crítica",
                 "Exposición ilustrativa (no aditiva) €"], result['suppliers'], [42, 19, 28, 19, 21, 35], (3, 6))
    evo = table_sheet(wb, "08_Evolucion", "EVOLUCIÓN MENSUAL / UNIVERSO 2026",
                      "Casos asignados al mes de la fecha de referencia del caso; no representa momento del posible error.",
                      ["Mes", "Facturas", "Casos", "Alta + crítica", "PO con sobreimporte", "Volumen absoluto facturas €"],
                      result['monthly'], [22, 18, 17, 22, 23, 28], (6,))
    ctl = table_sheet(wb, "09_Controles", "MATRIZ DE REGLAS DE DETECCIÓN",
                      "Las reglas son heurísticas de cribado, NO pruebas sustantivas concluyentes.",
                      ["Regla", "Nombre", "Puntos iniciales", "Prioridad teórica", "Condición",
                       "Casos (regla principal)", "Exposición orientativa €", "Familia"],
                      result['controls'], [14, 36, 21, 24, 84, 26, 30, 37], (7,))
    methods = [
        ["01", "Alcance", f"{len(result['src'].invoices)} facturas y {len(result['src'].orders)} POs sintéticas; CECOs y proyectos cargados para integridad referencial futura."],
        ["02", "R01", "Coincidencia de proveedor + número de factura normalizado. Caso crítico; revisar originales, abonos y registros."],
        ["03", "R02", "Misma huella proveedor + fecha + total firmado + descripción normalizada. No depende del número."],
        ["04", "R03", "Mismo proveedor, día e importe firmado. Riesgo elevado de falso positivo en proveedores de cargos repetitivos."],
        ["05", "R04", "Mismo proveedor e importe, fechas dentro de umbral y similitud de descripción >= configurada."],
        ["06", "R05", "Misma PO explícita e igual total con tolerancia. Facturación por fases puede justificar repetición."],
        ["07", "R06", "Facturado atribuible > PO, con tolerancia. Si una factura cita varias POs se prorratea sin evidencia externa."],
        ["08", "R07", "Proveedor PO distinto de proveedor factura; podría haber tercero facturador legítimo."],
        ["09", "Puntuación", "Puntuación base máxima por regla, +4 si comparte PO para indicios débiles, +2 si ambas facturas reportan pagado."],
        ["10", "Unidades", "Comparaciones de PO a bruto factura con IVA; no equivale a gasto del periodo ni pérdida por duplicado."],
        ["11", "Casos", "Los casos se identifican con IDs determinísticos. No sumar exposición: una factura puede repetir en varios pares."],
        ["12", "Confirmación", "Requiere factura PDF, proveedor, PO, contrato, recepción del servicio, mayor contable y evidencia bancaria."],
        ["13", "Abonos", "Se preserva el signo. Un abono negativo no se compara como duplicado de su factura positiva."],
        ["14", "Pendientes", "Estado 'Pagado' es dato del export ERP; este script NO cruza extractos bancarios."],
        ["15", "Seguimiento", "La hoja 02 incluye estados editables; para preservar revisión en re-ejecuciones importar seguimiento_casos.csv."],
        ["16", "Trazabilidad", "Fila origen se refiere a la hoja inicial del fichero X_DEMO correspondiente."],
        ["17", "Precaución", "Falsos positivos y negativos son esperables; no se dispone de PDF/adjuntos ni lógica de pagos reales."],
        ["18", "Privacidad", "Únicamente DEMO sintética SOCIEDAD FICTICIA X, S.L.; no publicar ficheros de terceros."],
    ]
    table_sheet(wb, "10_Metodologia", "METODOLOGÍA DE AUDITORÍA Y LIMITACIONES",
                "Documento de control, objetivo, afirmaciones de auditoría y procedimientos sugeridos en el manual.",
                ["ID", "Área", "Criterio metodológico / procedimiento"], methods, [13, 29, 115])
    ps = wb.create_sheet("11_Parametros")
    sheet_header(ps, "PARÁMETROS DEL MODELO", "La selección de mes se usa en dashboard; los umbrales se modifican con los argumentos CLI.", 5)
    for i, v in enumerate([["Mes", "Parámetro / valor"], ["Todos", ""]] + [[m, ""] for m in MONTHS], 5):
        ps.cell(row=i, column=1, value=v[0]); ps.cell(row=i, column=2, value=v[1] or None)
    for i, v in enumerate([["Parámetro", "Valor"], ["Días próximos", result['params']['days']],
                           ["Similitud mínima", result['params']['similarity']],
                           ["Tolerancia €", result['params']['threshold']], ["Motor versión", VERSION]], 5):
        ps.cell(row=i, column=4, value=v[0]); ps.cell(row=i, column=5, value=v[1])
    for ref in ("A5:B5", "D5:E5"):
        _style(ps, ref, BLUE, _font(10, True, WHITE))
    for cl in "ABCDE":
        ps.column_dimensions[cl].width = 22
    ps.freeze_panes = "A6"
    src_rows = []
    for category, name in result['src'].files.items():
        infile = (source_dir or (output.parents[1] / 'inputs')) / name
        digest = hashlib.sha256(infile.read_bytes()).hexdigest() if infile.exists() else 'Ruta no comprobable'
        count = (len(result['src'].invoices) if category == 'invoice' else len(result['src'].orders) if category == 'po'
                 else len(result['src'].project_ids) if category == 'project' else len(result['src'].cecos))
        src_rows.append([category, name, digest, count])
    table_sheet(wb, "12_Trazabilidad", "TRAZABILIDAD DE FUENTES SINTÉTICAS",
                "Hashes SHA-256 para verificar que se trabaja con el mismo conjunto de datos. NO conserva material original.",
                ["Conjunto", "Archivo de entrada", "SHA-256", "Registros leídos"], src_rows, [17, 55, 73, 20])

    # ---------------- dashboard ----------------
    dash = wb.create_sheet("01_Dashboard", 0)
    for k in "ABCDEFGHIJKLMNOP":
        dash.column_dimensions[k].width = 13
    dash.column_dimensions["A"].width = 17
    dash.column_dimensions["I"].width = 17
    sheet_header(dash, "DUPLICATE / FRAUD RISK DETECTOR  |  v" + VERSION.rsplit(".", 1)[0],
                 "SOCIEDAD FICTICIA X, S.L.  ·  Universo 2026  ·  Indicadores de revisión, NO fraude probado", 16)
    for x in (('FACTURAS REVISADAS', len(result['src'].invoices), 'A', 'D'), ('CASOS GENERADOS', len(cases), 'E', 'H'),
              ('CASOS PRIORITARIOS', high, 'I', 'L'), ('COINCIDENCIAS FUERTES', crit, 'M', 'P')):
        card(dash, *x)
    dash.row_dimensions[5].height = 25
    dash.merge_cells("A9:P9")
    _style(dash, "A9:P9", NAVY, _font(11, True, WHITE), Alignment(vertical="center", indent=1))
    dash["A9"] = "PANORAMA DE RIESGOS  /  CRITERIOS DEL PERIODO"
    dash.row_dimensions[9].height = 28
    for ref, val in (("A11:D11", "Pedidos analizados"), ("E11:F11", len(result['src'].orders)),
                     ("H11:K11", "Pares sospechosos"), ("L11:M11", pair),
                     ("A12:D12", "POs con sobreimporte"), ("E12:F12", po_cases),
                     ("H12:K12", "Facturas implicadas en pares"), ("L12:M12", result['invoice_case_count'])):
        dash.merge_cells(ref)
        dash[ref.split(":")[0]] = val
    dash.merge_cells("A15:B15"); dash["A15"] = "Filtrar mes:"
    dash.merge_cells("C15:E15"); dash["C15"] = "Todos"
    dv = DataValidation(type="list", formula1="'11_Parametros'!$A$6:$A$18", allow_blank=False)
    dash.add_data_validation(dv); dv.add("C15")
    _style(dash, "C15:E15", "#E0F4F1", _font(11, True, TEAL), Alignment(vertical="center", indent=1))
    dash.row_dimensions[15].height = 26
    dash.merge_cells("G15:J15"); dash["G15"] = "Casos en selección"
    dash.merge_cells("K15:L15")
    dash["K15"] = "=IF(C15=\"Todos\",SUM('08_Evolucion'!$C$6:$C$17),SUMIF('08_Evolucion'!$A$6:$A$17,C15,'08_Evolucion'!$C$6:$C$17))"
    dash.merge_cells("M15:O15"); dash["M15"] = "Alta + crítica"
    dash["P15"] = "=IF(C15=\"Todos\",SUM('08_Evolucion'!$D$6:$D$17),SUMIF('08_Evolucion'!$A$6:$A$17,C15,'08_Evolucion'!$D$6:$D$17))"
    dash.merge_cells("A18:P18"); dash["A18"] = "TENDENCIA MENSUAL Y DISTRIBUCIÓN DE CONTROLES"
    _style(dash, "A18:P18", INK, _font(11, True, WHITE), Alignment(vertical="center", indent=1))
    dash.row_dimensions[18].height = 27
    dash.merge_cells("A40:P40"); dash["A40"] = "ADVERTENCIAS DE INTERPRETACIÓN Y CONTROL DE CALIDAD"
    _style(dash, "A40:P40", INK, _font(11, True, WHITE), Alignment(vertical="center", indent=1))
    notes = ["• Cada caso es una hipótesis de revisión; no demuestra fraude ni pago duplicado.",
             "• Las cifras de exposición potencial NO se suman: una misma factura puede aparecer en varios casos.",
             "• 'Pagado' procede del fichero de compras; no está confirmado con bancos ni libro mayor.",
             "• R05 (misma PO e importe) puede ser facturación por hitos. R06 exige verificar suplementos y abonos.",
             "• Las fuentes son sintéticas y se conservan sin modificación. Revisión editable en 02_Casos."]
    for k, note in enumerate(notes, 42):
        dash.merge_cells(f"A{k}:P{k}")
        dash[f"A{k}"] = note
        _style(dash, f"A{k}:P{k}", PAPER if k % 2 == 0 else WHITE, _font(10, color=INK))
        dash.row_dimensions[k].height = 25
    cht = BarChart(); cht.type = "col"; cht.grouping = "clustered"
    cht.title = "Casos identificados por mes"
    cht.add_data(Reference(evo, min_col=3, min_row=5, max_row=17), titles_from_data=True)
    cht.set_categories(Reference(evo, min_col=1, min_row=6, max_row=17))
    cht.series[0].graphicalProperties.solidFill = BLUE.lstrip("#")
    cht.legend = None; cht.height = 9.5; cht.width = 25
    dash.add_chart(cht, "A20")
    cht2 = BarChart(); cht2.type = "bar"
    cht2.title = "Casos por regla principal"
    cht2.add_data(Reference(ctl, min_col=6, min_row=5, max_row=12), titles_from_data=True)
    cht2.set_categories(Reference(ctl, min_col=2, min_row=6, max_row=12))
    cht2.series[0].graphicalProperties.solidFill = TEAL.lstrip("#")
    cht2.legend = None; cht2.height = 9.5; cht2.width = 13
    dash.add_chart(cht2, "K20")
    dash.freeze_panes = "A5"
    dash.sheet_view.showGridLines = False
    wb.active = 0

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    stats = {"version": VERSION, "entity": ENTITY, "source_invoices": len(result['src'].invoices),
             "source_orders": len(result['src'].orders), "cases": len(cases), "duplicate_pairs": pair,
             "po_overbilling_cases": po_cases, "vendor_mismatch_cases": len(result['vendor_mismatch']),
             "critical": crit, "high_or_critical": high, "invoices_in_candidate_pairs": result['invoice_case_count'],
             "gross_invoice_total": total_invoice, "rules_count": len(RULES),
             "output_file": output.name, "engine": "openpyxl"}
    (output.parent / 'run_summary.json').write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding='utf-8')
    return stats


def main(argv=None):
    p=argparse.ArgumentParser(description="Duplicate & Fraud Risk Detector | Demo sintética")
    p.add_argument('--input-dir',type=Path,default=Path('inputs'))
    p.add_argument('--output',type=Path,default=Path('output/DUPLICATE_RISK_2026.xlsx'))
    p.add_argument('--review-file',type=Path,default=Path('inputs/seguimiento_casos.csv'))
    p.add_argument('--days',type=int,default=7)
    p.add_argument('--min-similarity',type=float,default=0.86)
    p.add_argument('--threshold',type=float,default=1.00)
    p.add_argument('--verbose',action='store_true')
    args=p.parse_args(argv)
    logging.basicConfig(level=logging.INFO if args.verbose else logging.WARNING,format='%(levelname)s: %(message)s')
    if args.days<0 or args.threshold<0 or not 0<=args.min_similarity<=1:p.error('Parámetros fuera de rango.')
    try:
        print('[1/3] Leyendo fuentes ficticias...')
        sources=load_data(args.input_dir)
        print('[2/3] Ejecutando controles con IDs estables...')
        result=detect(sources,args.days,args.min_similarity,args.threshold,read_review_csv(args.review_file))
        print('[3/3] Creando el Excel de auditoría...')
        summary=build_excel(result,args.output,args.input_dir)
        print(json.dumps(summary,indent=2,ensure_ascii=False))
    except Exception as err:
        logging.exception('No se pudo completar el análisis')
        raise SystemExit(f'ERROR: {err}') from err
    return summary

if __name__=='__main__':main()
