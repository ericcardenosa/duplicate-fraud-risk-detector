#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Apply the controlled test scenarios to the synthetic invoice export.

The base dataset (shared with the P2P Audit Engine) already contains R01
(repeated invoice numbers) and R06 (invoicing above the PO) cases. To make
rules R02-R05 visible, eight rows of the synthetic invoice file are rewritten
in place. Sheet names, headers and the number of rows are left unchanged, so
row numbers keep pointing to the same records.

The script is idempotent: every scenario sets fixed values, so running it on
an already prepared file produces the same file.

Usage:
    python scripts/create_demo_scenarios.py --file inputs/X_DEMO_Facturas_Compras.xlsx
"""
from __future__ import annotations

import argparse
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

# Column positions in the Holded-style purchase export (1-based).
COL = dict(date=1, num=2, acct_date=4, due=5, vendor=6, description=7, subtotal=12, iva=13,
           total=17, paid=18, outstanding=19, state=20, payment_date=21)
T = (10, 22, 34, 608000)  # time component used across the synthetic export


def d(y: int, m: int, day: int) -> datetime:
    return datetime(y, m, day, *T)


# (sheet row, rule illustrated, new values). Amounts carry 21 % VAT.
SCENARIOS = [
    (201, "R02", dict(date=d(2026, 5, 12), vendor="Proveedor Demo Control 901, S.L.",
                      description="Licencia de plataforma analitica / MAYO 2026", subtotal=1000.0, iva=210.0)),
    (202, "R02", dict(date=d(2026, 5, 12), vendor="Proveedor Demo Control 901, S.L.",
                      description="Licencia de plataforma analitica / MAYO 2026", subtotal=1000.0, iva=210.0)),
    (203, "R03", dict(date=d(2026, 6, 15), vendor="Proveedor Demo Control 902, S.L.",
                      description="Asistencia tecnica presencial - centro A", subtotal=2000.0, iva=420.0)),
    (205, "R03", dict(date=d(2026, 6, 15), vendor="Proveedor Demo Control 902, S.L.",
                      description="Material informatico extraordinario - centro B", subtotal=2000.0, iva=420.0)),
    (206, "R04", dict(date=d(2026, 8, 3), vendor="Proveedor Demo Control 903, S.L.",
                      description="Servicio de mantenimiento trimestral equipo central", subtotal=3000.0, iva=630.0)),
    (207, "R04", dict(date=d(2026, 8, 8), vendor="Proveedor Demo Control 903, S.L.",
                      description="Servicio mantenimiento trimestral equipo central", subtotal=3000.0, iva=630.0)),
    (672, "R05", dict(subtotal=5000.0, iva=1050.0)),   # same PO O2690023, milestone-style billing
    (1533, "R05", dict(subtotal=5000.0, iva=1050.0)),
]


def apply(path: Path) -> list[str]:
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    log = []
    for row, rule, v in SCENARIOS:
        num = ws.cell(row=row, column=COL["num"]).value
        if not (isinstance(num, str) and num.startswith("FX26-")):
            raise ValueError(f"La fila {row} no contiene una factura FX26-*: {num!r}")
        if "date" in v:
            for k in ("date", "acct_date"):
                ws.cell(row=row, column=COL[k]).value = v["date"]
            ws.cell(row=row, column=COL["due"]).value = v["date"] + timedelta(days=30)
        for k in ("vendor", "description", "subtotal", "iva"):
            if k in v:
                ws.cell(row=row, column=COL[k]).value = v[k]
        total = round(v["subtotal"] + v["iva"], 2)
        ws.cell(row=row, column=COL["total"]).value = total
        ws.cell(row=row, column=COL["paid"]).value = 0
        ws.cell(row=row, column=COL["outstanding"]).value = total
        ws.cell(row=row, column=COL["state"]).value = "Pendiente"
        ws.cell(row=row, column=COL["payment_date"]).value = None
        log.append(f"{rule}  fila {row:>4}  {num}  total {total:,.2f}")
    wb.save(path)
    return log


def main(argv=None):
    p = argparse.ArgumentParser(description="Aplica los escenarios de prueba R02-R05 a la copia demo de facturas.")
    p.add_argument("--file", type=Path, default=Path("inputs/X_DEMO_Facturas_Compras.xlsx"))
    args = p.parse_args(argv)
    for line in apply(args.file):
        print(line)


if __name__ == "__main__":
    main()
