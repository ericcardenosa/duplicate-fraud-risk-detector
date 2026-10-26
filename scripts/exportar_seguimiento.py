#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export the reviewer's work from 02_Casos to inputs/seguimiento_casos.csv.

Run it after editing the report and before regenerating it, so that status,
owner, evidence and conclusions are restored by case ID on the next run.

Usage:
    python scripts/exportar_seguimiento.py --report output/DUPLICATE_RISK_2026.xlsx --dest inputs/seguimiento_casos.csv
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

FIELDS = ["case_id", "estado", "responsable", "notas", "evidencia"]


def export(report: Path, dest: Path) -> int:
    wb = load_workbook(report, read_only=True, data_only=True)
    ws = wb["02_Casos"]
    header = [c for c in next(ws.iter_rows(min_row=5, max_row=5, values_only=True))]
    idx = {name: header.index(name) for name in
           ("ID caso", "Estado revisión", "Responsable", "Evidencia / enlace interno", "Conclusión / notas")}
    rows = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[idx["ID caso"]]:
            continue
        rec = {
            "case_id": r[idx["ID caso"]],
            "estado": r[idx["Estado revisión"]] or "Pendiente",
            "responsable": r[idx["Responsable"]] or "",
            "notas": r[idx["Conclusión / notas"]] or "",
            "evidencia": r[idx["Evidencia / enlace interno"]] or "",
        }
        if rec["estado"] != "Pendiente" or rec["responsable"] or rec["notas"] or rec["evidencia"]:
            rows.append(rec)
    wb.close()
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    return len(rows)


def main(argv=None):
    p = argparse.ArgumentParser(description="Exporta el seguimiento de 02_Casos a CSV.")
    p.add_argument("--report", type=Path, default=Path("output/DUPLICATE_RISK_2026.xlsx"))
    p.add_argument("--dest", type=Path, default=Path("inputs/seguimiento_casos.csv"))
    args = p.parse_args(argv)
    n = export(args.report, args.dest)
    print(f"{n} casos con revisión exportados a {args.dest}")


if __name__ == "__main__":
    main()
